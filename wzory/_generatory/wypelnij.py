# -*- coding: utf-8 -*-
"""Generuje folder 'Przyklad wypelniony' — wszystkie 13 wzorów na danych testowych.
Renderer uproszczony (na potrzeby podglądu); docelowy renderer po stronie aplikacji."""
import os, re, glob, shutil, copy
from docx import Document
import dane_testowe

WZORY = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CEL = os.path.join(WZORY, 'Przyklad wypelniony')
DANE = dane_testowe.DANE
SEKCJE = DANE['_sekcje']


def akapity(doc):
    for p in doc.paragraphs:
        yield p, doc
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    yield p, c


def rozwin_sekcje_inline(p):
    """Sekcja otwarta i domknięta w obrębie jednego akapitu."""
    tekst = p.text
    if '{{#' not in tekst:
        return
    wynik = tekst
    zmieniono = False
    for nazwa, elementy in SEKCJE.items():
        wzor = re.compile(r'\{\{#%s\}\}(.*?)\{\{/%s\}\}' % (nazwa, nazwa))
        if not wzor.search(wynik):
            continue
        zmieniono = True

        def zamien(m, el=elementy):
            out = []
            for e in el:
                frag = m.group(1)
                for k, v in e.items():
                    frag = frag.replace('{{%s}}' % k, str(v))
                out.append(frag)
            return ''.join(out)
        wynik = wzor.sub(zamien, wynik)
    if zmieniono:
        for i, r in enumerate(p.runs):
            r.text = wynik if i == 0 else ''


def podstaw_run(p, mapa):
    for r in p.runs:
        for k, v in mapa.items():
            if '{{%s}}' % k in r.text:
                r.text = r.text.replace('{{%s}}' % k, str(v))


def rozwin_sekcje_w_akapitach(doc):
    """Sekcje rozpięte na wielu akapitach: powiela akapity między znacznikami."""
    for nazwa, elementy in SEKCJE.items():
        while True:
            body = doc.element.body
            ps = doc.paragraphs
            i0 = i1 = None
            for i, p in enumerate(ps):
                if '{{#%s}}' % nazwa in p.text and i0 is None:
                    i0 = i
                elif '{{/%s}}' % nazwa in p.text and i0 is not None:
                    i1 = i; break
            if i0 is None or i1 is None:
                break
            szablon = [ps[j]._p for j in range(i0 + 1, i1)]
            kotwica = ps[i1]._p
            for el in elementy:
                for xp in szablon:
                    nowy = copy.deepcopy(xp)
                    kotwica.addprevious(nowy)
                    from docx.text.paragraph import Paragraph
                    podstaw_run(Paragraph(nowy, ps[i1]._parent), el)
            for xp in szablon:
                xp.getparent().remove(xp)
            ps[i0]._p.getparent().remove(ps[i0]._p)
            kotwica.getparent().remove(kotwica)


def rozwin_sekcje_w_wierszu(doc):
    """Sekcje w jednym wierszu tabeli: powiela wiersz."""
    for tb in doc.tables:
        for nazwa, elementy in SEKCJE.items():
            for row in list(tb.rows):
                tekst = ' '.join(c.text for c in row.cells)
                if '{{#%s}}' % nazwa not in tekst:
                    continue
                for el in elementy:
                    nowy = copy.deepcopy(row._tr)
                    row._tr.addprevious(nowy)
                    from docx.table import _Row
                    for c in _Row(nowy, tb).cells:
                        for p in c.paragraphs:
                            for r in p.runs:
                                r.text = (r.text.replace('{{#%s}}' % nazwa, '')
                                                .replace('{{/%s}}' % nazwa, ''))
                            podstaw_run(p, el)
                row._tr.getparent().remove(row._tr)


os.makedirs(CEL, exist_ok=True)

for f in sorted(glob.glob(os.path.join(WZORY, '*.docx'))):
    d = Document(f)
    rozwin_sekcje_w_wierszu(d)
    rozwin_sekcje_w_akapitach(d)
    for p, _ in akapity(d):
        rozwin_sekcje_inline(p)
        podstaw_run(p, {k: v for k, v in DANE.items() if k != '_sekcje'})
    nazwa = os.path.basename(f).replace('-WZOR.docx', '-TEST.docx')
    d.save(os.path.join(CEL, nazwa))
    print('ok', nazwa)
